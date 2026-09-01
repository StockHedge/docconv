"""스프레드시트 Writer: XLSX / CSV.

한글 깨짐 방지가 이 모듈의 핵심 관심사다.

* CSV 기본 인코딩은 ``utf-8-sig``(BOM 포함)다. Excel은 BOM이 없는 UTF-8 CSV를
  시스템 기본 코드페이지(한국어 Windows에서는 CP949)로 해석해 한글을 깨뜨린다.
  BOM 3바이트를 앞에 붙이면 Excel도 메모장도 UTF-8로 올바로 연다.
* XLSX는 내부가 UTF-8 XML이라 인코딩 문제가 없다.
"""

from __future__ import annotations

import csv
import datetime as _dt
from pathlib import Path
from typing import Any, Optional

from ..errors import WriteError
from ..ir import Document
from ..options import ConvertOptions, MergedCellPolicy
from ..util import units
from .flatten import Sheet, pick_primary_sheet, to_sheets


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------


def write_xlsx(doc: Document, path: Path | str, opts: ConvertOptions) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as e:  # pragma: no cover
        raise WriteError("openpyxl이 필요합니다. pip install openpyxl") from e

    sheets = to_sheets(doc, opts)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", start_color="FFF2F2F2")
    head_font = Font(bold=True)
    wrap = Alignment(vertical="top", wrap_text=True)

    for s in sheets:
        ws = wb.create_sheet(title=s.name)
        for ri, row in enumerate(s.rows, start=1):
            for ci, text in enumerate(row, start=1):
                src = s.cells[ri - 1][ci - 1] if ri - 1 < len(s.cells) else None
                value = _pick_value(text, src)
                c = ws.cell(row=ri, column=ci, value=value)
                if src is not None and src.number_format:
                    c.number_format = src.number_format
                if opts.keep_formatting:
                    c.alignment = wrap
                    if text:
                        c.border = border
                    if src is not None and src.background:
                        c.fill = PatternFill(
                            "solid",
                            start_color="FF"
                            + (units.rgb_to_hex(src.background, "") or "FFFFFF"),
                        )
            if s.header_row and ri == 1 and opts.keep_formatting:
                for ci in range(1, len(row) + 1):
                    ws.cell(row=1, column=ci).font = head_font
                    ws.cell(row=1, column=ci).fill = head_fill

        if opts.merged_cells is MergedCellPolicy.KEEP:
            for (r0, c0), (rs, cs) in s.merges.items():
                if rs > 1 or cs > 1:
                    try:
                        ws.merge_cells(
                            start_row=r0 + 1,
                            start_column=c0 + 1,
                            end_row=r0 + rs,
                            end_column=c0 + cs,
                        )
                    except ValueError:
                        pass

        _apply_widths(ws, s, get_column_letter)
        if s.header_row and s.rows:
            ws.freeze_panes = "A2"

    if not wb.sheetnames:
        wb.create_sheet(title="Sheet1")

    _set_props(wb, doc)
    try:
        wb.save(str(path))
    except OSError as e:
        raise WriteError(f"파일을 저장할 수 없습니다: {path}", detail=str(e)) from e


def _pick_value(text: str, src: Any) -> Any:
    """가능하면 원본 타입(숫자/날짜)을 살려 셀에 넣는다.

    문자열로만 쓰면 Excel에서 정렬/합계가 안 되고 왼쪽 정렬로 보인다.
    """
    if src is None:
        return text if text != "" else None
    raw = src.raw_value
    if raw is None:
        return text if text != "" else None
    if isinstance(raw, (int, float, bool, _dt.date, _dt.datetime, _dt.time)):
        return raw
    return text if text != "" else None


def _apply_widths(ws, s: Sheet, get_column_letter) -> None:
    n = s.n_cols
    if s.col_widths_pt:
        for i, w_pt in enumerate(s.col_widths_pt[:n], start=1):
            if w_pt and w_pt > 0:
                ws.column_dimensions[get_column_letter(i)].width = min(
                    120.0, units.pt_to_excel_col_width(w_pt)
                )
        return
    # 내용 길이로 추정한다. 한글은 폭이 2배에 가까우므로 가중치를 준다.
    for ci in range(n):
        longest = 0
        for row in s.rows[:400]:
            if ci < len(row):
                longest = max(longest, _display_width(row[ci]))
        if longest:
            ws.column_dimensions[get_column_letter(ci + 1)].width = min(
                60, max(8, longest + 2)
            )


def _display_width(s: str) -> int:
    first_line = s.split("\n")[0] if s else ""
    return sum(2 if ord(ch) > 0x1100 else 1 for ch in first_line)


def _set_props(wb, doc: Document) -> None:
    p = wb.properties
    if doc.meta.title:
        p.title = doc.meta.title
    if doc.meta.author:
        p.creator = doc.meta.author
    if doc.meta.subject:
        p.subject = doc.meta.subject
    if doc.meta.keywords:
        p.keywords = doc.meta.keywords


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


def write_csv(doc: Document, path: Path | str, opts: ConvertOptions) -> list[Path]:
    """CSV로 쓴다. 시트가 여러 장이면 파일을 나눠 쓴다.

    반환값은 실제로 만들어진 파일 목록이다(첫 항목이 주 출력물).
    """
    sheets = to_sheets(doc, opts)
    base = Path(path)
    made: list[Path] = []

    if len(sheets) <= 1:
        _write_one_csv(sheets[0] if sheets else Sheet("Sheet1", [], []), base, opts)
        return [base]

    # 어느 시트를 주 출력(이름.csv)으로 삼을지 결정한다. 첫 시트가 늘 알맹이는
    # 아니다 - 배경은 flatten.pick_primary_sheet 의 주석 참고.
    primary = pick_primary_sheet(sheets)
    order = [primary] + [i for i in range(len(sheets)) if i != primary]

    for rank, idx in enumerate(order):
        s = sheets[idx]
        target = (
            base
            if rank == 0
            else base.with_name(f"{base.stem}_{_safe_stem(s.name)}{base.suffix}")
        )
        _write_one_csv(s, target, opts)
        made.append(target)
    return made


def _write_one_csv(s: Sheet, path: Path, opts: ConvertOptions) -> None:
    try:
        with path.open(
            "w", encoding=opts.csv_encoding, newline="", errors="replace"
        ) as f:
            w = csv.writer(
                f,
                delimiter=opts.csv_delimiter,
                lineterminator=opts.csv_line_terminator,
                quoting=csv.QUOTE_MINIMAL,
            )
            for row in s.rows:
                w.writerow([_csv_cell(v) for v in row])
    except OSError as e:
        raise WriteError(f"파일을 저장할 수 없습니다: {path}", detail=str(e)) from e


def _csv_cell(v: str) -> str:
    """셀 안의 줄바꿈은 유지하되(따옴표로 감싸짐) CR은 제거해 이중 개행을 막는다."""
    return v.replace("\r\n", "\n").replace("\r", "\n") if v else ""


def _safe_stem(name: str) -> str:
    bad = set('\\/:*?"<>|')
    return "".join("_" if ch in bad else ch for ch in name).strip() or "sheet"
