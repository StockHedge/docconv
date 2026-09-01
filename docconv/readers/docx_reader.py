"""DOCX(Office Open XML) Reader.

python-docx를 쓰되, `document.paragraphs` 와 `document.tables` 를 따로 읽으면
**문서 내 순서가 사라진다**. 표가 전부 뒤로 밀리거나 문단 사이 위치를 잃는다.
그래서 `body` 의 XML 자식을 직접 순회해 원래 순서를 보존한다.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterator, Optional

from ..errors import CorruptFileError, ReadError
from ..ir import (
    Align,
    Block,
    Cell,
    DocKind,
    DocMeta,
    Document,
    Image,
    ListKind,
    PageBreak,
    Paragraph,
    Run,
    Table,
)
from ..util import units

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _q(tag: str) -> str:
    return "{%s}%s" % (W_NS, tag)


_ALIGN_MAP = {
    "left": Align.LEFT,
    "start": Align.LEFT,
    "center": Align.CENTER,
    "right": Align.RIGHT,
    "end": Align.RIGHT,
    "both": Align.JUSTIFY,
    "justify": Align.JUSTIFY,
    "distribute": Align.DISTRIBUTE,
}


class DocxReader:
    def __init__(self, path: Path | str) -> None:
        self.path = Path(path)
        self.warnings: list[str] = []

    def read(self) -> Document:
        try:
            import docx  # python-docx
            from docx.document import Document as _DocxDoc
        except ImportError as e:  # pragma: no cover
            raise ReadError(
                "python-docx가 설치되어 있지 않습니다. pip install python-docx"
            ) from e

        try:
            d = docx.Document(str(self.path))
        except Exception as e:
            raise CorruptFileError(
                "DOCX 파일을 열 수 없습니다.", detail=f"{self.path}: {e}"
            ) from e

        self._doc = d
        doc = Document(kind=DocKind.DOCUMENT, meta=self._meta(d))
        self._apply_section(d, doc)

        from docx.table import Table as _T
        from docx.text.paragraph import Paragraph as _P

        body = d.element.body
        for child in body.iterchildren():
            tag = child.tag
            if tag == _q("p"):
                blocks = list(self._paragraph(_P(child, d)))
                doc.blocks.extend(blocks)
            elif tag == _q("tbl"):
                doc.blocks.append(self._table(_T(child, d)))
            elif tag == _q("sectPr"):
                continue

        doc.compact()
        if not doc.meta.title:
            doc.meta.title = _first_text(doc)
        return doc

    # -- 메타/구역 --------------------------------------------------------

    @staticmethod
    def _meta(d) -> DocMeta:
        cp = d.core_properties
        return DocMeta(
            title=cp.title or None,
            author=cp.author or None,
            subject=cp.subject or None,
            keywords=cp.keywords or None,
            created=cp.created.isoformat() if cp.created else None,
            modified=cp.modified.isoformat() if cp.modified else None,
            source_format="docx",
        )

    @staticmethod
    def _apply_section(d, doc: Document) -> None:
        try:
            s = d.sections[0]
        except (IndexError, AttributeError):
            return
        if s.page_width and s.page_height:
            doc.page_width_pt = s.page_width.pt
            doc.page_height_pt = s.page_height.pt
        doc.margin_pt = (
            s.left_margin.pt if s.left_margin else 56.7,
            s.top_margin.pt if s.top_margin else 56.7,
            s.right_margin.pt if s.right_margin else 56.7,
            s.bottom_margin.pt if s.bottom_margin else 56.7,
        )

    # -- 문단 -------------------------------------------------------------

    def _paragraph(self, p) -> Iterator[Block]:
        style_name = None
        try:
            style_name = p.style.name if p.style is not None else None
        except Exception:
            pass

        para = Paragraph(style_name=style_name)
        para.heading = _heading_level(style_name)
        para.align = _ALIGN_MAP.get(
            str(p.alignment).split(".")[-1].split(" ")[0].lower()
            if p.alignment is not None
            else "",
            Align.LEFT,
        )

        pf = p.paragraph_format
        try:
            if pf.left_indent:
                para.indent_pt = pf.left_indent.pt
            if pf.space_before:
                para.space_before_pt = pf.space_before.pt
            if pf.space_after:
                para.space_after_pt = pf.space_after.pt
            if pf.line_spacing and isinstance(pf.line_spacing, float):
                para.line_spacing = pf.line_spacing
        except Exception:
            pass

        # 목록 여부는 numPr 존재로 판단한다.
        if p._p.find(".//" + _q("numPr")) is not None:
            para.list_kind = ListKind.BULLET
            ilvl = p._p.find(".//" + _q("ilvl"))
            if ilvl is not None:
                para.list_level = int(ilvl.get(_q("val")) or 0)
        elif style_name and "List" in style_name:
            para.list_kind = ListKind.BULLET

        pending: list[Block] = []
        for r in p.runs:
            # 페이지 나누기
            if r._r.find(".//" + _q("br") + "[@" + _q("type") + "='page']") is not None:
                pending.append(PageBreak())
            img = self._run_image(r)
            if img is not None:
                pending.append(img)
            if r.text:
                para.runs.append(self._run(r))

        para.merge_runs()
        if para.runs:
            yield para
        yield from pending

    @staticmethod
    def _run(r) -> Run:
        f = r.font
        color = None
        try:
            if f.color is not None and f.color.rgb is not None:
                s = str(f.color.rgb)
                color = units.hex_to_rgb(s)
        except Exception:
            pass
        highlight = None
        try:
            if f.highlight_color is not None:
                highlight = (255, 255, 0)
        except Exception:
            pass
        return Run(
            text=r.text,
            bold=bool(f.bold),
            italic=bool(f.italic),
            underline=bool(f.underline),
            strike=bool(f.strike),
            superscript=bool(f.superscript),
            subscript=bool(f.subscript),
            size_pt=f.size.pt if f.size else None,
            font=f.name,
            color=color,
            highlight=highlight,
        )

    def _run_image(self, r) -> Optional[Image]:
        blips = r._r.findall(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
        )
        if not blips:
            return None
        embed = blips[0].get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        if not embed:
            return None
        try:
            part = self._doc.part.related_parts[embed]
            data = part.blob
            ext = Path(str(part.partname)).suffix.lower().lstrip(".")
        except Exception:
            return None
        w = h = None
        ext_el = r._r.find(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/main}ext"
        )
        if ext_el is not None:
            w = units.emu_to_pt(int(ext_el.get("cx") or 0)) or None
            h = units.emu_to_pt(int(ext_el.get("cy") or 0)) or None
        return Image(
            data=data,
            fmt="jpeg" if ext in ("jpg", "jpeg") else (ext or "png"),
            width_pt=w,
            height_pt=h,
            name=Path(str(part.partname)).name,
        )

    # -- 표 ---------------------------------------------------------------

    def _table(self, t) -> Table:
        """표를 XML 레벨에서 직접 읽는다.

        python-docx의 ``row.cells`` 는 병합 셀을 span 수만큼 반복해서 돌려주므로
        중복을 걸러야 하는데, ``id(tc._tc)`` 로 거르면 안 된다. lxml은 요소에
        접근할 때마다 새 프록시를 만들고, 이전 프록시가 GC되면 **같은 메모리
        주소가 재사용**된다. 그러면 서로 다른 셀이 같은 id를 갖게 되어 멀쩡한
        셀이 병합 자리로 오판되고 내용이 통째로 사라진다(실측: 12,061자 ->
        6,093자).

        그래서 ``<w:tr>`` 의 ``<w:tc>`` 를 직접 순회하고 ``gridSpan`` / ``vMerge``
        를 읽어 그리드를 구성한다.
        """
        from docx.table import _Cell

        tbl = t._tbl
        trs = tbl.findall(_q("tr"))
        if not trs:
            return Table()

        grid: list[list[Optional[Cell]]] = []
        #: 열 위치 -> 수직 병합을 시작한 셀 (row_span을 늘려가기 위해)
        open_vmerge: dict[int, Cell] = {}

        for tr in trs:
            row_cells: list[Optional[Cell]] = []
            col = 0
            for tc in tr.findall(_q("tc")):
                span = _grid_span(tc)
                vmerge = _vmerge(tc)

                if vmerge == "continue":
                    anchor = open_vmerge.get(col)
                    if anchor is not None:
                        anchor.row_span += 1
                    for _ in range(span):
                        row_cells.append(Cell(merged_placeholder=True))
                    col += span
                    continue

                c = Cell()
                c.col_span = span
                c.blocks = list(self._cell_blocks(_Cell(tc, t)))
                c.background = _cell_shading(tc)
                row_cells.append(c)
                # 가로 병합으로 가려지는 칸을 placeholder로 채운다.
                for _ in range(span - 1):
                    row_cells.append(Cell(merged_placeholder=True))

                if vmerge == "restart":
                    open_vmerge[col] = c
                else:
                    open_vmerge.pop(col, None)
                col += span
            grid.append(row_cells)

        out = Table(rows=[[c if c is not None else Cell() for c in r] for r in grid])
        out.normalize()
        out.header_row = _looks_like_header(t)
        out.col_widths_pt = _col_widths(t)
        return out

    def _cell_blocks(self, tc) -> Iterator[Block]:
        from docx.table import Table as _T
        from docx.text.paragraph import Paragraph as _P

        for child in tc._tc.iterchildren():
            if child.tag == _q("p"):
                yield from self._paragraph(_P(child, tc))
            elif child.tag == _q("tbl"):
                yield self._table(_T(child, tc))


# --------------------------------------------------------------------------
# 보조
# --------------------------------------------------------------------------


def _grid_span(tc) -> int:
    """가로 병합 칸 수.

    ``.//`` 로 찾으면 **중첩 표 안쪽 셀의 gridSpan**까지 잡혀 바깥 표의 그리드가
    어긋난다. 반드시 이 셀의 tcPr만 본다.
    """
    gs = tc.find(_q("tcPr") + "/" + _q("gridSpan"))
    if gs is None:
        return 1
    try:
        return max(1, int(gs.get(_q("val")) or 1))
    except (ValueError, TypeError):
        return 1


def _vmerge(tc) -> Optional[str]:
    """세로 병합 상태. 'restart' | 'continue' | None.

    ``<w:vMerge/>`` 처럼 val 속성이 없으면 'continue'가 기본값이다.
    """
    vm = tc.find(_q("tcPr") + "/" + _q("vMerge"))
    if vm is None:
        return None
    val = (vm.get(_q("val")) or "continue").lower()
    return "restart" if val == "restart" else "continue"


def _cell_shading(tc) -> Optional[tuple[int, int, int]]:
    # gridSpan과 같은 이유로 이 셀의 tcPr만 본다(중첩 표 오염 방지).
    shd = tc.find(_q("tcPr") + "/" + _q("shd"))
    if shd is None:
        return None
    fill = shd.get(_q("fill"))
    if not fill or fill.lower() in ("auto", "ffffff"):
        return None
    return units.hex_to_rgb(fill)


def _col_widths(t) -> Optional[list[float]]:
    grid = t._tbl.find(_q("tblGrid"))
    if grid is None:
        return None
    out: list[float] = []
    for gc in grid.findall(_q("gridCol")):
        w = gc.get(_q("w"))
        if w:
            try:
                out.append(units.twip_to_pt(int(w)))
            except ValueError:
                pass
    return out or None


def _looks_like_header(t) -> bool:
    """첫 행에 tblHeader가 있거나 모든 셀이 굵으면 머리글로 본다."""
    try:
        first = t.rows[0]
    except IndexError:
        return False
    if first._tr.find(".//" + _q("tblHeader")) is not None:
        return True
    texts = [c.text.strip() for c in first.cells]
    if not any(texts):
        return False
    for c in first.cells:
        for p in c.paragraphs:
            for r in p.runs:
                if r.text.strip() and not r.font.bold:
                    return False
    return True


_HEADING_PREFIXES = ("heading ", "제목 ", "개요 ")


def _heading_level(style_name: Optional[str]) -> int:
    if not style_name:
        return 0
    low = style_name.strip().lower()
    if low in ("title", "제목"):
        return 1
    if low in ("subtitle", "부제"):
        return 2
    for pre in _HEADING_PREFIXES:
        if low.startswith(pre):
            try:
                return max(1, min(9, int(low[len(pre) :].strip())))
            except ValueError:
                return 0
    return 0


def _first_text(doc: Document) -> Optional[str]:
    for b in doc.blocks:
        if isinstance(b, Paragraph) and not b.is_blank():
            return b.text.strip()[:120]
        if isinstance(b, Table):
            for row in b.rows:
                for c in row:
                    if c.text.strip():
                        return c.text.strip()[:120]
    return None


def read(path: Path | str) -> Document:
    r = DocxReader(path)
    return _attach(r.read(), r)


def _attach(doc, reader):
    """Reader가 모은 경고를 Document에 실어 파이프라인까지 전달한다."""
    if getattr(reader, "warnings", None):
        doc.meta.extra.setdefault("warnings", []).extend(reader.warnings)
    return doc
